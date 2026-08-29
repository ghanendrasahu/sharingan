"""DOCX text extraction service."""

from __future__ import annotations

from dataclasses import dataclass, field


@dataclass
class DOCXParagraph:
    """A paragraph from a DOCX document."""

    index: int
    text: str
    style: str | None = None
    is_heading: bool = False
    heading_level: int | None = None


@dataclass
class DOCXExtractionResult:
    """Result of DOCX text extraction."""

    paragraphs: list[DOCXParagraph] = field(default_factory=list)
    full_text: str = ""
    total_paragraphs: int = 0
    total_chars: int = 0
    total_words: int = 0
    metadata: dict = field(default_factory=dict)
    error: str | None = None

    def __post_init__(self):
        self.total_paragraphs = len(self.paragraphs)
        self.full_text = "\n\n".join(p.text for p in self.paragraphs if p.text.strip())
        self.total_chars = len(self.full_text)
        self.total_words = len(self.full_text.split()) if self.full_text else 0


def extract_docx_text(docx_data: bytes) -> DOCXExtractionResult:
    """Extract text from DOCX using python-docx."""
    try:
        from io import BytesIO

        from docx import Document

        doc = Document(BytesIO(docx_data))
        paragraphs = []

        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                # Detect heading style
                is_heading = para.style and para.style.name.startswith("Heading")
                heading_level = None
                if is_heading and para.style.name:
                    import re

                    match = re.search(r"Heading\s*(\d)", para.style.name)
                    if match:
                        heading_level = int(match.group(1))

                paragraphs.append(
                    DOCXParagraph(
                        index=i,
                        text=para.text,
                        style=para.style.name if para.style else None,
                        is_heading=is_heading,
                        heading_level=heading_level,
                    )
                )

        # Extract metadata
        metadata = {}
        if doc.core_properties:
            props = doc.core_properties
            metadata = {
                "title": props.title or "",
                "author": props.author or "",
                "subject": props.subject or "",
                "created": str(props.created) if props.created else "",
                "modified": str(props.modified) if props.modified else "",
            }

        return DOCXExtractionResult(
            paragraphs=paragraphs,
            metadata=metadata,
        )
    except ImportError:
        return DOCXExtractionResult(
            error="python-docx not installed. Install with: pip install python-docx"
        )
    except Exception as e:
        return DOCXExtractionResult(error=f"DOCX extraction failed: {e!s}")


def extract_docx_text_fallback(docx_data: bytes) -> DOCXExtractionResult:
    """Basic DOCX extraction using ZIP and XML parsing."""
    try:
        import re
        from io import BytesIO
        from zipfile import ZipFile

        with ZipFile(BytesIO(docx_data)) as zf:
            # Read document.xml
            with zf.open("word/document.xml") as f:
                content = f.read().decode("utf-8", errors="ignore")

            # Extract text from XML tags
            text_parts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", content)

            # Create paragraphs from text
            paragraphs = []
            for i, text in enumerate(text_parts):
                if text.strip():
                    paragraphs.append(
                        DOCXParagraph(
                            index=i,
                            text=text,
                        )
                    )

            return DOCXExtractionResult(paragraphs=paragraphs)
    except Exception as e:
        return DOCXExtractionResult(error=f"Fallback DOCX extraction failed: {e!s}")
